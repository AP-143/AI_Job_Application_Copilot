"""Extract raw text from an uploaded CV file (PDF, DOCX, or plain text)."""
from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedFileTypeError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Route to the right parser based on file extension and return plain text."""
    lower = filename.lower()

    if lower.endswith(".pdf"):
        text = _extract_pdf_text(content)
    elif lower.endswith(".docx"):
        text = _extract_docx_text(content)
    elif lower.endswith(".txt"):
        text = content.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type for '{filename}'. Use PDF, DOCX, or TXT."
        )

    text = text.strip()
    if not text:
        raise EmptyDocumentError(
            "No extractable text found. If this is a scanned/image CV, try exporting a text-based PDF."
        )
    return text


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts: list[str] = [p.text for p in doc.paragraphs]

    # Tables (some CVs use them for layout, e.g. skills grids)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts)
